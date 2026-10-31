# -*- coding: utf-8 -*-
import csv
import io
import os
import re
import socket
import threading
import time
import asyncio
import webbrowser
from urllib.parse import urlparse

from flask import Flask, jsonify, render_template_string, request

try:
    from telethon import TelegramClient, errors
    from telethon.sessions import StringSession
    from telethon.tl.types import User
    TELETHON_OK = True
except Exception:
    TELETHON_OK = False

app = Flask(__name__)

# ====================== НАСТРОЙКИ ======================
TG_API_ID = os.environ.get("TG_API_ID")
TG_API_HASH = os.environ.get("TG_API_HASH")
SEND_DELAY = int(os.environ.get("SEND_DELAY", "60"))

# ====================== ФИЛЬТРЫ ======================
NOT_A_SITE = {
    "vk.com", "vk.ru", "m.vk.com", "instagram.com", "facebook.com", "fb.com", "ok.ru",
    "t.me", "telegram.me", "telegram.dog", "wa.me", "whatsapp.com", "api.whatsapp.com",
    "max.ru", "youtube.com", "youtu.be", "tiktok.com", "twitter.com", "x.com",
    "taplink.cc", "taplink.ru", "linktr.ee", "yandex.ru", "ya.ru", "2gis.ru",
    "zoon.ru", "profi.ru", "avito.ru", "dikidi.net", "dikidi.ru", "yclients.com",
}
TG_SKIP = {"joinchat", "s", "addstickers", "share", "c", "proxy", "socks", "login", "iv", "setlanguage"}
TG_USER_RE = re.compile(r"^[A-Za-z][A-Za-z0-9_]{3,31}$")
TOKEN_SPLIT_RE = re.compile(r"[\s,;]+")

COLS = {
    "id": ["id организации", "id"], "name": ["название", "name"], "sites": ["сайты", "сайт", "website", "site"],
    "social": ["соцсети", "соцсеть", "social"], "yandex": ["ссылка на карточку", "ссылка яндекс", "yandex"],
    "category": ["основная категория", "категория"], "city": ["город"],
}

def host_of(token: str) -> str:
    t = token.strip()
    if not re.match(r"^[a-z]+://", t, re.I): t = "http://" + t
    try: h = (urlparse(t).hostname or "").lower()
    except ValueError: return ""
    return h[4:] if h.startswith("www.") else h

def is_real_site(cell: str) -> bool:
    for tok in TOKEN_SPLIT_RE.split(cell or ""):
        if not tok or "." not in tok: continue
        h = host_of(tok)
        if h and h not in NOT_A_SITE and not any(h.endswith("." + d) for d in NOT_A_SITE): return True
    return False

def extract_contacts(*cells):
    found, seen = [], set()
    for cell in cells:
        for tok in TOKEN_SPLIT_RE.split(cell or ""):
            if not tok: continue
            tok = tok.strip("()<>\"'")
            h = host_of(tok)
            if h in ("t.me", "telegram.me", "telegram.dog"):
                t = tok if re.match(r"^[a-z]+://", tok, re.I) else "https://" + tok
                path = [p for p in urlparse(t).path.split("/") if p]
                if not path: continue
                user = path[0]
                if user.startswith("+") or user.lower() in TG_SKIP or not TG_USER_RE.match(user): continue
                url = f"https://t.me/{user}"
                if url.lower() not in seen:
                    seen.add(url.lower()); found.append({"type": "tg", "url": url})
            elif h == "max.ru":
                t = tok if re.match(r"^[a-z]+://", tok, re.I) else "https://" + tok
                p = urlparse(t)
                path = [x for x in p.path.split("/") if x]
                if not path or path[0].lower() == "join" or path[0].lstrip("+").isdigit(): continue
                url = f"https://max.ru{p.path}".rstrip("/")
                if url.lower() not in seen:
                    seen.add(url.lower()); found.append({"type": "max", "url": url})
    return found

def read_rows(filename: str, data: bytes):
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        rows = [["" if v is None else str(v) for v in r] for r in ws.iter_rows(values_only=True)]
    else:
        text = data.decode("utf-8-sig", errors="ignore")
        try: delim = csv.Sniffer().sniff(text[:5000], delimiters=";,\t").delimiter
        except: delim = ";"
        rows = list(csv.reader(io.StringIO(text, newline=""), delimiter=delim))
    return rows[0], rows[1:]

def find_col(header, keys):
    low = [h.strip().lower() for h in header]
    for k in keys:
        if k in low: return low.index(k)
    return None

def process(filename: str, data: bytes):
    header, body = read_rows(filename, data)
    idx = {k: find_col(header, v) for k, v in COLS.items()}
    if idx["name"] is None: raise ValueError("Не нашёл колонку «Название»")
    def get(row, key):
        i = idx[key]
        return row[i].strip() if i is not None and i < len(row) else ""

    stats = {"total": 0, "with_site": 0, "no_contact": 0, "duplicates": 0}
    items, seen_ids = [], set()
    for row in body:
        name = get(row, "name")
        if not name: continue
        stats["total"] += 1
        oid = get(row, "id")
        if oid and oid in seen_ids: stats["duplicates"] += 1; continue
        seen_ids.add(oid)
        if is_real_site(get(row, "sites")): stats["with_site"] += 1; continue
        contacts = extract_contacts(get(row, "social"), get(row, "sites"))
        if not contacts: stats["no_contact"] += 1; continue
        yandex = get(row, "yandex") or (f"https://yandex.ru/maps/org/{oid}" if oid else "")
        items.append({"id": oid, "name": name, "contacts": contacts, "yandex": yandex, "category": get(row, "category"), "city": get(row, "city")})
    stats["selected"] = len(items)
    return {"items": items, "stats": stats}

# ====================== СЕНДЕР ======================
SEND_STATE = {
    "running": False, "paused": False, "stop_requested": False,
    "current_index": -1, "total": 0, "sent": 0, "failed": 0, "skipped": 0,
    "log": [], "sent_ids": set(), "skipped_ids": set(), "manual_done": set(),
}
_sender_thread = None
_state_lock = threading.RLock()

def _reset_state(total=0, manual_done=()):
    with _state_lock:
        SEND_STATE.update({
            "running": True, "paused": False, "stop_requested": False,
            "current_index": -1, "total": total, "sent": 0, "failed": 0, "skipped": 0,
            "log": [], "sent_ids": set(), "skipped_ids": set(),
            "manual_done": set(manual_done),
        })

def _should_stop():
    with _state_lock: return SEND_STATE["stop_requested"]

def _wait_if_paused():
    while True:
        with _state_lock: paused = SEND_STATE["paused"]
        if not paused or _should_stop(): break
        time.sleep(0.5)

def _log(text):
    with _state_lock:
        SEND_STATE["log"].append(text)
        if len(SEND_STATE["log"]) > 200: SEND_STATE["log"] = SEND_STATE["log"][-200:]

async def _send_one(client, item, message_text):
    tg_contact = next((c for c in item["contacts"] if c["type"] == "tg"), None)
    if not tg_contact: return False, "no_tg_contact"
    username = tg_contact["url"].rstrip("/").split("/")[-1]
    try:
        entity = await client.get_entity(username)
    except (errors.UsernameNotOccupiedError, errors.UsernameInvalidError, ValueError):
        return False, "username_not_found"
    except errors.FloodWaitError as e:
        return False, f"flood_wait_{e.seconds}s"
    except Exception as e:
        return False, f"entity_error: {type(e).__name__}: {str(e)[:80]}"
    if not isinstance(entity, User): return False, f"not_a_user ({type(entity).__name__})"
    if getattr(entity, "bot", False): return False, "is_bot"

    # уже есть переписка с этим человеком (мы писали или он нам) — не пишем
    try:
        hist = await client.get_messages(entity, limit=1)
        if len(hist) > 0: return False, "already_chat"
    except errors.FloodWaitError as e:
        return False, f"flood_wait_{e.seconds}s"
    except Exception as e:
        # если историю проверить не вышло — на всякий случай НЕ пишем, чтобы не продублировать
        return False, f"history_error: {type(e).__name__}"

    try:
        await client.send_message(entity, message_text)
        return True, "sent"
    except errors.FloodWaitError as e:
        return False, f"flood_wait_{e.seconds}s"
    except errors.UserPrivacyRestrictedError:
        return False, "privacy_restricted"
    except Exception as e:
        return False, f"send_error: {type(e).__name__}"

STRIKE_PREFIXES = ("not_a_user", "is_bot", "username_not_found")

def _sender_worker(items, tpl):
    loop = asyncio.new_event_loop()
    asyncio.set_event_loop(loop)

    async def _run():
        if not TELETHON_OK: _log("Telethon не установлен"); return
        session_str = os.environ.get("TG_SESSION_STRING", "").strip()
        if not session_str: _log("Нет сессии. Зайди в /setup."); return

        client = TelegramClient(StringSession(session_str), int(TG_API_ID), TG_API_HASH)
        try:
            await client.connect()
            if not await client.is_user_authorized(): _log("Сессия не авторизована."); return

            for idx, item in enumerate(items):
                if _should_stop(): _log("Остановлено."); break
                _wait_if_paused()
                if _should_stop(): break

                key = item["id"] or item["name"]
                with _state_lock:
                    SEND_STATE["current_index"] = idx
                    already = key in SEND_STATE["manual_done"]

                # 1) Ты сам поставил галочку «Написал(а)» — пропускаем сразу, без кд
                if already:
                    with _state_lock: SEND_STATE["skipped"] += 1
                    _log(f"☑ {item['name']} — уже отмечен, пропуск")
                    continue

                msg = tpl.replace("{name}", item["name"]).replace("{yandex}", item["yandex"] or "")
                ok, reason = await _send_one(client, item, msg)

                with _state_lock:
                    if ok:
                        SEND_STATE["sent"] += 1
                        SEND_STATE["sent_ids"].add(key)
                        _log(f"✅ {item['name']} — отправлено")
                    elif reason == "already_chat":
                        # уже есть переписка — не пишем, без кд, помечаем как «Написал(а)»
                        SEND_STATE["skipped"] += 1
                        SEND_STATE["sent_ids"].add(key)
                        _log(f"💬 {item['name']} — уже есть переписка (пропущено)")
                    elif reason == "no_tg_contact":
                        # только MAX — не зачёркиваем, просто идём дальше без кд
                        SEND_STATE["skipped"] += 1
                        _log(f"⏭ {item['name']} — только MAX (пропущено)")
                    elif reason.startswith(STRIKE_PREFIXES):
                        # канал / группа / бот / аккаунт не найден — просто идём дальше, без кд и без зачёркивания
                        SEND_STATE["skipped"] += 1
                        _log(f"⏭ {item['name']} — {reason} (пропущено)")
                    else:
                        SEND_STATE["failed"] += 1
                        _log(f"❌ {item['name']} — {reason}")

                if not ok: await asyncio.sleep(1)

                # 3) Кулдаун ТОЛЬКО если реально отправили сообщение
                if ok and idx < len(items) - 1:
                    waited = 0
                    while waited < SEND_DELAY and not _should_stop():
                        _wait_if_paused(); time.sleep(0.5); waited += 0.5
        finally:
            await client.disconnect()

    try: loop.run_until_complete(_run())
    except Exception as e: _log(f"Критическая ошибка: {e}")
    finally:
        loop.close()
        with _state_lock:
            SEND_STATE["running"] = False
            SEND_STATE["paused"] = False

@app.post("/send_start")
def send_start():
    global _sender_thread
    data = request.get_json(silent=True) or {}
    items = data.get("items")
    tpl = data.get("tpl", "")
    if not items: return jsonify(error="Нет списка"), 400
    _reset_state(total=len(items), manual_done=data.get("done", []))
    _sender_thread = threading.Thread(target=_sender_worker, args=(items, tpl), daemon=True)
    _sender_thread.start()
    return jsonify(ok=True)

@app.post("/mark_done")
def mark_done():
    data = request.get_json(silent=True) or {}
    key = data.get("key")
    if not key: return jsonify(error="no key"), 400
    with _state_lock:
        if data.get("done"): SEND_STATE["manual_done"].add(key)
        else: SEND_STATE["manual_done"].discard(key)
    return jsonify(ok=True)

@app.post("/send_pause")
def send_pause():
    with _state_lock:
        if not SEND_STATE["running"]: return jsonify(error="Не запущено"), 400
        SEND_STATE["paused"] = True
    return jsonify(ok=True)

@app.post("/send_resume")
def send_resume():
    with _state_lock:
        if not SEND_STATE["running"]: return jsonify(error="Не запущено"), 400
        SEND_STATE["paused"] = False
    return jsonify(ok=True)

@app.post("/send_stop")
def send_stop():
    with _state_lock:
        SEND_STATE["stop_requested"] = True; SEND_STATE["paused"] = False
    return jsonify(ok=True)

@app.get("/send_status")
def send_status():
    with _state_lock:
        return jsonify({
            "running": SEND_STATE["running"], "paused": SEND_STATE["paused"],
            "current_index": SEND_STATE["current_index"], "total": SEND_STATE["total"],
            "sent": SEND_STATE["sent"], "failed": SEND_STATE["failed"], "skipped": SEND_STATE["skipped"],
            "sent_ids": list(SEND_STATE["sent_ids"]),
            "skipped_ids": list(SEND_STATE["skipped_ids"]),
            "log": SEND_STATE["log"][-30:],
        })

@app.post("/upload")
def upload():
    f = request.files.get("file")
    if not f: return jsonify(error="Файл не выбран"), 400
    try: return jsonify(process(f.filename or "", f.read()))
    except Exception as e: return jsonify(error=str(e)), 400

@app.get("/")
def index(): return render_template_string(PAGE)

# ====================== HTML ======================
PAGE = r"""<!doctype html>
<html lang="ru"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>Рассылка</title>
<style>
:root{--bg:#f6f7f9;--card:#fff;--tx:#1b1f24;--mut:#6b7480;--bd:#e2e5ea;--ac:#2f6fed;--ok:#1f9d55;--warn:#e6a817;--danger:#d33}
@media(prefers-color-scheme:dark){:root{--bg:#14171a;--card:#1d2125;--tx:#e8eaed;--mut:#9aa3ad;--bd:#2d3339;--ac:#6b9bff;--ok:#4cc27f;--warn:#f0b400;--danger:#ff6b6b}}
*{box-sizing:border-box}body{margin:0;font:15px/1.5 system-ui,sans-serif;background:var(--bg);color:var(--tx)}
.wrap{max-width:960px;margin:0 auto;padding:24px 16px 80px}
h1{font-size:22px;margin:0 0 16px}.card{background:var(--card);border:1px solid var(--bd);border-radius:12px;padding:16px;margin-bottom:16px}
#drop{border:2px dashed var(--bd);border-radius:12px;padding:28px;text-align:center;cursor:pointer;color:var(--mut)}
#drop.over{border-color:var(--ac);color:var(--ac)}
textarea{width:100%;font:inherit;color:inherit;background:transparent;border:1px solid var(--bd);border-radius:8px;padding:10px}
button{border:1px solid var(--bd);background:var(--card);border-radius:8px;padding:7px 12px;cursor:pointer;}
button.pri{background:var(--ac);color:#fff;border-color:var(--ac)}button.ok{background:var(--ok);color:#fff;border-color:var(--ok)}
button.warn{background:var(--warn);color:#000;border-color:var(--warn)}button.danger{background:var(--danger);color:#fff;border-color:var(--danger)}
button:disabled{opacity:.5;cursor:not-allowed}
.row{display:flex;gap:8px;flex-wrap:wrap;align-items:center}.mut{color:var(--mut);font-size:13px}
.item.done{opacity:.55}.item h3{margin:0 0 4px;font-size:16px}
.msg{margin:10px 0;white-space:pre-wrap;background:var(--bg);border-radius:8px;padding:10px;font-size:14px}
#send-log{max-height:200px;overflow:auto;font-size:13px;background:var(--bg);border-radius:8px;padding:8px;margin-top:8px}
.badge{display:inline-block;font-size:12px;padding:1px 7px;border-radius:99px;background:var(--bd);color:var(--mut)}
.badge.run{background:#e6f0ff;color:#2f6fed}.badge.pause{background:#fff3cd;color:#7a5c00}
</style></head><body><div class="wrap">
<h1>Бизнесы без сайта → готовые сообщения</h1>
<div class="card"><div id="drop">Перетащи сюда файл <b>.csv</b> или <b>.xlsx</b></div><input id="file" type="file" hidden><div id="err" style="color:var(--danger);margin-top:8px"></div></div>
<div class="card"><b>Шаблон сообщения</b> <span class="mut">(подставляется {name})</span><textarea id="tpl" rows="7"></textarea></div>

<div id="result" hidden>
  <div class="card row" style="justify-content:space-between">
    <div id="stats"></div>
    <label><input type="checkbox" id="hideDone"> Скрыть отправленные</label>
  </div>
  <div class="card">
    <div class="row" style="justify-content:space-between;margin-bottom:10px"><b>Автоматическая отправка в ТГ</b> <span id="send-badge" class="badge">не запущено</span></div>
    <div class="row">
      <button id="btnStart" class="pri">Отправить в Telegram</button>
      <button id="btnPause" class="warn" disabled>Пауза</button>
      <button id="btnResume" class="ok" disabled>Продолжить</button>
      <button id="btnStop" class="danger" disabled>Остановить</button>
      <span class="mut" id="send-progress"></span>
    </div>
    <div id="send-log"></div>
  </div>
  <div id="list"></div>
</div>
</div>
<script>
const DEFAULT_TPL = `Здравствуйте! Я Арсений, заметил, что у «{name}» пока нет сайта — а сейчас многие клиенты перед обращением гуглят или смотрят Яндекс.Карты. Без сайта часть из них просто уходит к конкурентам, у которых он есть.
Перед визитом или звонком людям важно изучить точные цены, почитать отзывы и убедиться в качестве работы. По статистике, наличие удобной веб-страницы увеличивает количество обращений из карт в среднем на 30–40%.
Я 4 года создаю сайты под ключ для локального бизнеса. Делаю их быстро, качественно и без лишней головной боли для вас: полностью адаптирую под мобильные телефоны, оформляю понятный прайс, блок с отзывами и контакты.
Было бы интересно сделать для вас сайт? Могу предоставить вам свою последнюю работу для примера`;
const $ = id => document.getElementById(id);
let items = [];
let done = new Set(JSON.parse(localStorage.getItem('done') || '[]'));
localStorage.removeItem('skipped');
$('tpl').value = localStorage.getItem('tpl_v2') || DEFAULT_TPL;

function msgFor(it) { return $('tpl').value.replaceAll('{name}', it.name); }

function render(){
  const list = $('list'); list.innerHTML = '';
  items.forEach((it, i) => {
    const key = it.id || it.name;
    const isDone = done.has(key);
    if ($('hideDone').checked && isDone) return;
    const d = document.createElement('div'); d.className = 'card item' + (isDone ? ' done' : '');
    d.innerHTML = `<h3>${i+1}. ${it.name}</h3>`;
    const links = document.createElement('div'); links.className = 'row'; links.style.margin = '8px 0';
    it.contacts.forEach(c => { const a = document.createElement('a'); a.href = c.url; a.target = '_blank'; a.textContent = (c.type === 'tg' ? 'Telegram' : 'MAX'); links.append(a); });
    if (it.yandex) { const y = document.createElement('a'); y.href = it.yandex; y.target = '_blank'; y.textContent = 'Яндекс Карты'; links.append(y); }
    d.append(links);
    const m = document.createElement('div'); m.className = 'msg'; m.textContent = msgFor(it); d.append(m);
    const act = document.createElement('div'); act.className = 'row';
    const cp = document.createElement('button'); cp.textContent = 'Копировать'; cp.onclick = () => { navigator.clipboard.writeText(msgFor(it)); cp.textContent = '✓'; setTimeout(()=>cp.textContent='Копировать',1000); };
    act.append(cp);
    const lb = document.createElement('label'); lb.style.marginLeft = 'auto';
    const cb = document.createElement('input'); cb.type = 'checkbox'; cb.checked = isDone;
    cb.onchange = () => {
      cb.checked ? done.add(key) : done.delete(key);
      localStorage.setItem('done', JSON.stringify([...done]));
      fetch('/mark_done', {method:'POST', headers:{'Content-Type':'application/json'}, body: JSON.stringify({key, done: cb.checked})});
      render();
    };
    lb.append(cb, ' Написал(а)'); act.append(lb);
    d.append(act); list.append(d);
  });
}

$('drop').onclick = () => $('file').click();
$('file').onchange = e => {
  const f = e.target.files[0]; if(!f) return;
  const fd = new FormData(); fd.append('file', f);
  $('err').textContent = 'Загрузка...';
  fetch('/upload', {method:'POST', body: fd}).then(r=>r.json()).then(j => {
    if (j.error) { $('err').textContent = j.error; return; }
    items = j.items; $('err').textContent = '';
    $('stats').innerHTML = `<span>Всего: <b>${j.stats.total}</b></span> · <span>Подходит: <b>${j.stats.selected}</b></span>`;
    $('result').hidden = false; render();
  }).catch(e => $('err').textContent = 'Ошибка: ' + e);
};
['dragover','dragenter'].forEach(ev => $('drop').addEventListener(ev, e => { e.preventDefault(); $('drop').classList.add('over'); }));
['dragleave','drop'].forEach(ev => $('drop').addEventListener(ev, e => { e.preventDefault(); $('drop').classList.remove('over'); }));
$('drop').addEventListener('drop', e => { $('file').files = e.dataTransfer.files; $('file').onchange({target:{files:e.dataTransfer.files}}); });

$('tpl').oninput = () => { localStorage.setItem('tpl_v2', $('tpl').value); render(); };
$('hideDone').onchange = render;

function setSendUI(st){
  const running = st.running, paused = st.paused;
  $('btnStart').disabled = running; $('btnPause').disabled = !running || paused; $('btnResume').disabled = !running || !paused; $('btnStop').disabled = !running;
  const badge = $('send-badge'); badge.className = 'badge' + (running ? (paused ? ' pause' : ' run') : ''); badge.textContent = running ? (paused ? 'пауза' : 'работает') : 'не запущено';
  $('send-progress').textContent = running ? `Обработано: ${st.current_index+1} / ${st.total} · отправлено ${st.sent} · ошибок ${st.failed} · пропущено ${st.skipped}` : '';
  $('send-log').innerHTML = (st.log || []).map(x => `<div>${x}</div>`).join('');
  let changed = false;
  (st.sent_ids || []).forEach(id => { if (!done.has(id)) { done.add(id); changed = true; } });
  if (changed){
    localStorage.setItem('done', JSON.stringify([...done]));
    render();
  }
}
function pollSend(){ fetch('/send_status').then(r => r.json()).then(setSendUI).catch(()=>{}); }
setInterval(pollSend, 2000); pollSend();

$('btnStart').onclick = () => {
  if (!items.length){ alert('Сначала загрузи файл'); return; }
  const todo = items.filter(it => !done.has(it.id || it.name));
  if (!todo.length){ alert('Все уже отмечены как написанные'); return; }
  fetch('/send_start', {method:'POST', headers:{'Content-Type':'application/json'},
    body: JSON.stringify({items: todo, tpl: $('tpl').value, done: [...done]})
  }).then(r => r.json()).then(j => { if (j.error) alert(j.error); pollSend(); });
};
$('btnPause').onclick = () => fetch('/send_pause', {method:'POST'}).then(pollSend);
$('btnResume').onclick = () => fetch('/send_resume', {method:'POST'}).then(pollSend);
$('btnStop').onclick = () => { if(confirm('Остановить рассылку?')) fetch('/send_stop', {method:'POST'}).then(pollSend); };
</script></body></html>"""

if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5000)
